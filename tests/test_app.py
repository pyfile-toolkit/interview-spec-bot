"""Тесты интервью-бота: путь до PDF, границы валидации, доступ к артефактам."""

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import store
from app.main import app
from app.questions import QUESTIONS

client = TestClient(app)

# Ответы, которые проходят валидацию по длине и содержат цифры (чтобы «что уточнить» был осмысленным).
FULL_ANSWERS = {
    "product": "Телеграм-бот для записи клиентов в барбершоп, с выбором мастера и времени",
    "audience": "Клиенты барбершопа, 300 человек в месяц; решает владелец салона",
    "features": "Запись на слот; напоминание за 2 часа; выгрузка записей в Excel по дням",
    "done": "Открываю бота, жму /start, выбираю мастера и время — запись появляется у администратора",
    "deadline": "До 20 октября, привязано к открытию второго салона",
    "budget": "Оплата 45 000 рублей по этапам, поддержка 3 месяца после сдачи",
}


@pytest.fixture(autouse=True)
def clean_db():
    store.reset()
    yield
    store.reset()


def new_session() -> dict:
    r = client.post("/api/session")
    assert r.status_code == 200
    return r.json()


def answer_all(session_id: str, answers=None) -> dict:
    data = answers or FULL_ANSWERS
    state = None
    for q in QUESTIONS:
        r = client.post("/api/answer", json={"session_id": session_id, "answer": data[q.id]})
        assert r.status_code == 200, r.text
        state = r.json()
    return state


def test_index_renders_chat_page():
    r = client.get("/")
    assert r.status_code == 200
    assert "Интервью → ТЗ" in r.text
    assert "api/answer" in r.text


def test_new_session_starts_with_first_question():
    st = new_session()
    assert st["complete"] is False
    assert st["question"]["id"] == QUESTIONS[0].id
    assert st["answers_count"] == 0
    assert st["total_questions"] == len(QUESTIONS)


def test_full_interview_produces_pdf_and_xlsx():
    st = new_session()
    final = answer_all(st["session_id"])
    assert final["complete"] is True
    assert final["question"] is None
    assert final["pdf_url"].endswith(".pdf")

    pdf = client.get(final["pdf_url"])
    assert pdf.status_code == 200
    assert pdf.content[:4] == b"%PDF"
    assert len(pdf.content) > 1500

    xlsx = client.get(final["xlsx_url"])
    assert xlsx.status_code == 200
    ws = load_workbook(io.BytesIO(xlsx.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Вопрос", "Ответ")
    assert len(rows) - 1 == len(QUESTIONS)
    assert rows[1][1] == FULL_ANSWERS["product"]


def test_empty_answer_rejected_422():
    sid = new_session()["session_id"]
    r = client.post("/api/answer", json={"session_id": sid, "answer": "   "})
    assert r.status_code == 422
    assert "Пустой ответ" in r.json()["detail"]


def test_too_short_answer_rejected_422():
    sid = new_session()["session_id"]
    q = QUESTIONS[0]
    r = client.post("/api/answer", json={"session_id": sid, "answer": "бот"})
    assert r.status_code == 422
    assert str(q.min_len) in r.json()["detail"]


def test_pdf_before_finish_is_409():
    sid = new_session()["session_id"]
    client.post("/api/answer", json={"session_id": sid, "answer": FULL_ANSWERS["product"]})
    r = client.get(f"/api/spec/{sid}.pdf")
    assert r.status_code == 409
    assert client.get(f"/api/spec/{sid}.xlsx").status_code == 409


def test_unknown_session_is_404():
    assert client.get("/api/session/deadbeef").status_code == 404
    assert client.post("/api/answer", json={"session_id": "deadbeef", "answer": "x" * 40}).status_code == 404
    assert client.get("/api/spec/deadbeef.pdf").status_code == 404
    assert client.get("/api/spec/deadbeef.xlsx").status_code == 404


def test_answer_after_finish_is_409():
    sid = new_session()["session_id"]
    answer_all(sid)
    r = client.post("/api/answer", json={"session_id": sid, "answer": "ещё одна мысль про бюджет"})
    assert r.status_code == 409


def test_clarify_section_flags_vague_budget():
    """Ответ без цифры в бюджете => PDF просит зафиксировать сумму (не выдумываем, а помечаем)."""
    from app.pdf import _clarify_points

    points = _clarify_points({**FULL_ANSWERS, "budget": "оплатим как договоримся"})
    assert any("Бюджет" in p for p in points)
    assert not any("Бюджет" in p for p in _clarify_points(FULL_ANSWERS))
